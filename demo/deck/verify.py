"""Validate source bindings, derived values, editable deck structure, PDFs and optional real export (HACK-036, v2 HACK-059)."""
from __future__ import annotations
import argparse
from collections import Counter
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import re
import unicodedata

import pymupdf
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / 'demo/deck'


def tokens(text):
    return Counter(re.findall(r'[\w.+%-]+', unicodedata.normalize('NFKC', text)))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--rendered', type=Path, help='PDF exported from the real PPTX with LibreOffice')
    args = parser.parse_args()
    sources = json.loads((DECK/'sources.json').read_text(encoding='utf-8'))
    assert sources['demo_data'] is False and sources['exploratory'] is True
    for path, expected in sources['source_sha256'].items():
        assert hashlib.sha256((ROOT/path).read_bytes()).hexdigest() == expected, f'Stale source: {path}'
    prs = Presentation(DECK/'deck.pptx')
    assert len(prs.slides) == 10
    pdf = pymupdf.open(DECK/'deck.pdf')
    one = pymupdf.open(ROOT/'demo/one-pager/one-pager.pdf')
    assert len(pdf) == 10 and len(one) == 1
    for entry in sources['evidence']:
        rows = list(csv.DictReader((ROOT/entry['source']).open(encoding='utf-8')))
        selected = [r for r in rows if all(r[k] == str(v) for k,v in entry['row'].items())]
        assert len(selected) == 1
        actual = float(selected[0][entry['column']])
        assert actual == entry['value']
        assert entry['display'] in pdf[entry['slide']-1].get_text(), entry
    for entry in sources['derived']:
        vals = []
        for i in entry['inputs']:
            rows = list(csv.DictReader((ROOT/i['source']).open(encoding='utf-8')))
            sel = [r for r in rows if all(r[k] == str(v) for k, v in i['row'].items())]
            assert len(sel) == 1, i
            vals.append(Decimal(sel[0][i['column']]))
        v = vals[0] - vals[1] if entry['formula'] == 'a-b' else (1 - vals[0] / vals[1]) * 100
        shown = str(v.quantize(Decimal('1').scaleb(-entry['digits']), rounding=ROUND_HALF_UP)).replace('-', '−')
        assert shown == entry['display'], entry
        assert entry['display'].lstrip('−') in pdf[entry['slide']-1].get_text(), entry
    all_text = '\n'.join(page.get_text() for page in pdf)
    assert 'No se suman' in pdf[6].get_text() and 'Identificación inferida' in pdf[6].get_text()
    assert 'sin confirmarse' in pdf[0].get_text() and 'Identificación inferida' in pdf[0].get_text()
    assert 'escenario' in pdf[8].get_text()
    assert 'Stuff+ v2' in all_text and 'No mejora' not in all_text
    assert 'Identificación inferida' in one[0].get_text()
    # A geometry check alone does not replace rendering; source text and real
    # export must also survive, and pages are inspected visually by the author.
    for i, slide in enumerate(prs.slides):
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left+shape.width <= prs.slide_width
            assert shape.top+shape.height <= prs.slide_height
        editable = ' '.join(s.text for s in slide.shapes if s.has_text_frame)
        assert tokens(editable) == tokens(pdf[i].get_text()), f'PPTX/PDF content mismatch on {i+1}'
    for document in (pdf,one):
        for page in document:
            assert '\ufffd' not in page.get_text()
            for block in page.get_text('dict')['blocks']:
                if 'lines' not in block: continue
                for line in block['lines']:
                    for span in line['spans']:
                        x0,y0,x1,y1 = span['bbox']
                        assert x0 >= 0 and y0 >= 0 and x1 <= page.rect.width+.5 and y1 <= page.rect.height+.5
    if args.rendered:
        real = pymupdf.open(args.rendered)
        assert len(real) == 10
        for i,page in enumerate(real):
            assert tokens(page.get_text()) == tokens(pdf[i].get_text()), f'Rendered PPTX lost text on {i+1}'
    print(f'PASS: 10 editable slides; 10-page presentation PDF; 1-page summary; {len(sources["evidence"])} verified CSV values; {len(sources["derived"])} derived values; source hashes current' + ('; real PPTX export text matches' if args.rendered else ''))


if __name__ == '__main__': main()
