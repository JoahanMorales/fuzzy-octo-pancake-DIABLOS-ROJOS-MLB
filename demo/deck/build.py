"""Build the editable PPTX + presentation PDF + one-page PDF from tracked aggregates (HACK-036, v2 in HACK-059).

    python demo/deck/build.py

Every number on a slide is read from a CSV under reports/ (file, row, column and SHA-256 go to sources.json);
the two derived numbers (league average and bias reduction) record their inputs and formula, and verify.py
recomputes them. One layout (class Page) drives three outputs: the PPTX, a ReportLab PDF with the bundled DejaVu
font, and a LaTeX source (Beamer for the deck, article for the one-pager, TikZ at the same coordinates, Fira Sans)
that Tectonic compiles over the ReportLab PDFs when it is installed (HACK-088). Same lines, same words, so
verify.py's PPTX/PDF text check holds for either PDF. No private data is read.
"""
from __future__ import annotations
import csv
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import shutil
import subprocess

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'demo/deck'
ONE = ROOT / 'demo/one-pager'
FONTS = OUT / 'fonts'
ASSETS = OUT / 'assets'
RED, INK, GRAY, LIGHT, WHITE = '#D71920', '#202124', '#585B60', '#ECEDEF', '#FFFFFF'
BG, PANEL, TEXT, MUTED = '#121316', '#1F2125', '#F4F4F5', '#9EA1A7'  # dark slides
W, H = 720, 405
PPTX_FONT = 'Arial'  # what PowerPoint shows when editing; the PDF embeds the bundled font
for name, file in [('Deck', 'DejaVuSans.ttf'), ('DeckBold', 'DejaVuSans-Bold.ttf')]:
    pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))

sources, evidence, derived = {}, [], []


def row(path, **where):
    p = ROOT / path
    # LF line endings, as git stores the file: a Windows checkout has CRLF, so the hash must not depend on the OS
    sources[path] = hashlib.sha256(p.read_bytes().replace(b'\r\n', b'\n')).hexdigest()
    rows = list(csv.DictReader(p.open(encoding='utf-8')))
    found = [r for r in rows if all(r[k] == str(v) for k, v in where.items())]
    if len(found) != 1:
        raise ValueError(f'Expected one source row: {path}, {where}, found {len(found)}')
    return found[0]


def fmt(value: Decimal, digits: int) -> str:
    s = str(value.quantize(Decimal('1').scaleb(-digits), rounding=ROUND_HALF_UP))
    return s.replace('-', '−')


def val(display: str) -> float:
    return float(display.replace("−", "-"))


def number(slide, path, where, column, digits=2, factor=1):
    r = row(path, **where)
    display = fmt(Decimal(r[column]) * Decimal(str(factor)), digits)
    evidence.append(dict(slide=slide, source=path, row=where, column=column,
                         value=float(r[column]), factor=factor, digits=digits, display=display))
    return display


def derive(slide, formula, inputs, digits):
    """formula 'a-b' or '1-a/b' (as a percentage) on CSV inputs [(path, where, column), ...]."""
    vals = [Decimal(row(p, **w)[c]) for p, w, c in inputs]
    if formula == 'a-b':
        v = vals[0] - vals[1]
    elif formula == '1-a/b':
        v = (1 - vals[0] / vals[1]) * 100
    else:
        raise ValueError(formula)
    display = fmt(v, digits)
    derived.append(dict(slide=slide, formula=formula, digits=digits, display=display,
                        inputs=[dict(source=p, row=w, column=c) for p, w, c in inputs]))
    return display


TEX_ESCAPE = {'\\': r'\textbackslash{}', '%': r'\%', '&': r'\&', '#': r'\#', '_': r'\_', '$': r'\$',
              '{': r'\{', '}': r'\}', '~': r'\textasciitilde{}', '^': r'\textasciicircum{}'}


def tex_escape(s: str) -> str:
    return ''.join(TEX_ESCAPE.get(ch, ch) for ch in s)


def tex_color(hexcolor: str) -> str:
    r, g, b = (int(hexcolor[i:i+2], 16) for i in (1, 3, 5))
    return f'{{rgb,255:red,{r};green,{g};blue,{b}}}'


class Page:
    """Shared explicit positions for native PowerPoint shapes, the ReportLab PDF and the LaTeX source.

    tex is a list that collects TikZ commands for this page (coordinates in bp from the bottom left)."""
    def __init__(self, ppt=None, pdf=None, width=W, height=H, tex=None, tex_base=OUT):
        self.ppt, self.pdf, self.width, self.height = ppt, pdf, width, height
        self.tex, self.tex_base = tex, tex_base

    def rect(self, x, y, w, h, color, radius=0):
        assert x >= 0 and y >= 0 and x+w <= self.width+.1 and y+h <= self.height+.1
        if self.ppt is not None:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
            shape = self.ppt.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
            if radius: shape.adjustments[0] = min(.5, radius / max(min(w, h), 1))
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color[1:])
            shape.line.fill.background()
            style = shape._element.find(qn("p:style"))
            if style is not None: shape._element.remove(style)
        if self.pdf is not None:
            self.pdf.setFillColor(color)
            if radius: self.pdf.roundRect(x, self.height-y-h, w, h, radius, stroke=0, fill=1)
            else: self.pdf.rect(x, self.height-y-h, w, h, stroke=0, fill=1)
        if self.tex is not None:
            corners = f', rounded corners={radius}bp' if radius else ''
            self.tex.append(f'\\fill[color={tex_color(color)}{corners}] ({x:.2f},{self.height-y:.2f}) '
                            f'rectangle ({x+w:.2f},{self.height-y-h:.2f});')

    def image(self, path, x, y, w, h):
        assert x >= 0 and y >= 0 and x+w <= self.width+.1 and y+h <= self.height+.1
        if self.ppt is not None:
            self.ppt.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))
        if self.pdf is not None:
            self.pdf.drawImage(str(path), x, self.height-y-h, w, h)
        if self.tex is not None:
            rel = Path(path).resolve().relative_to(self.tex_base).as_posix()
            self.tex.append(f'\\node[anchor=north west, inner sep=0] at ({x:.2f},{self.height-y:.2f}) '
                            f'{{\\includegraphics[width={w}bp,height={h}bp]{{{rel}}}}};')

    def text(self, text, x, y, width=640, size=18, color=INK, bold=False, leading=None):
        font = 'DeckBold' if bold else 'Deck'
        leading = leading or size*1.24
        lines = []
        for paragraph in str(text).split('\n'):
            line = ''
            for word in paragraph.split():
                candidate = (line+' '+word).strip()
                if pdfmetrics.stringWidth(candidate, font, size) > width and line:
                    lines.append(line); line = word
                else: line = candidate
            lines.append(line)
        for i, line in enumerate(lines):
            yy = y+i*leading
            assert yy+size <= self.height-8, (line, yy, self.height)
            assert pdfmetrics.stringWidth(line, font, size) <= width+.1, line
            if self.ppt is not None:
                shape = self.ppt.shapes.add_textbox(Pt(x), Pt(yy-2), Pt(width), Pt(leading+5))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]; p.text = line
                p.font.name = PPTX_FONT; p.font.size = Pt(size); p.font.bold = bold
                p.font.color.rgb = RGBColor.from_string(color[1:])
            if self.pdf is not None:
                self.pdf.setFillColor(color); self.pdf.setFont(font, size)
                self.pdf.drawString(x, self.height-yy-size*.82, line)
            if self.tex is not None and line:
                series = '\\bfseries' if bold else ''
                self.tex.append(f'\\node[anchor=base west, inner sep=0, text={tex_color(color)}, '
                                f'font=\\fontsize{{{size}bp}}{{{size*1.2:.1f}bp}}\\selectfont{series}] '
                                f'at ({x:.2f},{self.height-yy-size*.82:.2f}) {{{tex_escape(line)}}};')
        return y+len(lines)*leading


TEX_HEAD = {
    'deck': r"""\documentclass[t]{beamer}
\geometry{papersize={720bp,405bp}}
\setbeamertemplate{navigation symbols}{}
\hypersetup{pdftitle={Stuff+ 2240 · Diablos Rojos · deck v2}}
""",
    'one-pager': r"""\documentclass{article}
\usepackage[paperwidth=595bp,paperheight=842bp,margin=0pt]{geometry}
\pagestyle{empty}
\usepackage{hyperref}
\hypersetup{pdftitle={Stuff+ 2240 · resumen ejecutivo v2}}
""",
}
TEX_COMMON = r"""\usepackage[sfdefault,lining]{FiraSans}
\usepackage{graphicx}
\usepackage{tikz}
\setlength{\parindent}{0pt}
\begin{document}
"""


def tex_document(kind: str, pages: list[list[str]]) -> str:
    """Pages of TikZ commands at absolute positions (bp from the bottom left of the page)."""
    out = ['% Generado por demo/deck/build.py (HACK-088). No editar: cambiar build.py y reconstruir.',
           TEX_HEAD[kind] + TEX_COMMON]
    for cmds in pages:
        pic = ('\\begin{tikzpicture}[remember picture, overlay, x=1bp, y=1bp, shift={(current page.south west)}]\n'
               + '\n'.join(cmds) + '\n\\end{tikzpicture}')
        out.append(f'\\begin{{frame}}[plain]\n{pic}\n\\end{{frame}}' if kind == 'deck' else f'\\null\n{pic}\n\\newpage')
    out.append('\\end{document}\n')
    return '\n'.join(out)


def typeset(deck_pages, one_pages):
    """Write deck.tex and one-pager.tex; with Tectonic installed, compile them over the ReportLab PDFs."""
    jobs = [(OUT/'deck.tex', tex_document('deck', deck_pages)), (ONE/'one-pager.tex', tex_document('one-pager', [one_pages]))]
    for path, text in jobs:
        path.write_text(text, encoding='utf-8', newline='\n')
    tectonic = shutil.which('tectonic')
    if not tectonic:
        print('tectonic no está instalado: se quedan los PDF de ReportLab (DejaVu Sans).')
        return
    for path, _ in jobs:
        run = subprocess.run([tectonic, '-X', 'compile', path.name], cwd=path.parent, capture_output=True, text=True)
        if run.returncode != 0:
            raise RuntimeError(f'tectonic falló con {path.name}:\n{run.stderr[-3000:]}')


def header(p, title, index, source):
    p.rect(0, 0, W, H, BG)
    p.rect(0, 0, 6, H, RED)
    p.rect(40, 24, 8, 8, RED)
    p.text('STUFF+ 2240   /   DIABLOS ROJOS', 56, 21, 480, 9, MUTED, True)
    p.text(f'{index:02d}', 656, 21, 30, 11, RED, True)
    p.text(title, 40, 48, 640, 25, TEXT, True, 30)
    p.rect(40, 360, 640, .8, PANEL)
    p.text(source, 40, 369, 610, 7.5, MUTED, leading=10)


def card(p, x, y, w, h):
    p.rect(x, y, w, h, PANEL, 6)


def stat(p, x, y, w, value, label, color=RED, size=40, label_size=13):
    """A big number and its label inside a card that starts at (x, y)."""
    p.text(value, x+16, y+16, w-32, size, color, True)
    p.text(label, x+16, y+24+size*1.15, w-32, label_size, TEXT, leading=label_size*1.3)


def bar(p, label, value, text, y, max_value, color=RED, label_w=165):
    p.text(label, 40, y, label_w, 14, TEXT)
    p.rect(215, y, 335, 18, PANEL, 3)
    p.rect(215, y, max(abs(value)/max_value*335, 6), 18, color, 3)
    p.text(text, 564, y, 116, 16, TEXT, True)


def main():
    OUT.mkdir(parents=True, exist_ok=True); ONE.mkdir(parents=True, exist_ok=True)
    prs = Presentation(); prs.slide_width = Pt(W); prs.slide_height = Pt(H)
    pdf = Canvas(str(OUT/'deck.pdf'), pagesize=(W, H))
    pdf.setTitle('Stuff+ 2240 · Diablos Rojos · deck v2')
    titles, footers, tex_pages = [], [], []

    def slide(title, footer, note):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.notes_slide.notes_text_frame.text = note + '\nFuentes: ' + footer
        tex_pages.append([])
        p = Page(s, pdf, tex=tex_pages[-1]); header(p, title, len(prs.slides), footer)
        titles.append(title); footers.append(footer)
        return p

    def end(): pdf.showPage()

    shape = 'reports/altitude_delta/delta_by_shape.csv'
    league = derive(1, 'a-b', [(shape, {'shape': 'sweeper'}, 'delta_stuff'),
                               (shape, {'shape': 'sweeper'}, 'relative_to_league')], 1)

    # 1 · the headline
    p = slide(f'El mismo arsenal pierde {league.replace("−", "")} puntos\nde Stuff+ en el aire de CDMX',
              shape + ' · promedio de la liga = ΔStuff+ de una forma − su diferencia contra la liga.',
              '45 s. Problema: un buen arsenal a nivel del mar vale menos en casa. Comparamos el mismo lanzamiento en dos aires con Stuff+ v2 en producción. Abrir con la decisión del coach: qué reclutar y cómo usarlo en casa.')
    p.text(league, 40, 128, 300, 84, RED, True)
    p.text('ΔStuff+ de la liga en el aire del Harp Helú.', 345, 140, 335, 14.5, TEXT, True, 20)
    p.text('El modelo mide el aire de cada juego y lo extrapola a un aire que no vio.', 345, 190, 335, 14, MUTED, leading=19)
    p.rect(40, 262, 640, 58, PANEL, 6)
    p.rect(40, 262, 4, 58, RED)
    p.text('Datos oficiales agregados. Identificación inferida: la sede del staff\nsigue sin confirmarse.', 58, 274, 610, 13, TEXT, leading=18)
    end()

    # 2 · the air
    density = 'reports/physics/rho_by_bucket.csv'
    p = slide('El aire estimado cambia entre ambientes', density + ' · media por bucket; parques anónimos.',
              '55 s. Densidad relativa estimada con el arrastre de las rectas. Son grupos de altitud; la identidad exacta del Harp Helú no está verificada.')
    for y, name, label in [(140, 'No Altitude', 'Nivel del mar'), (196, 'Medium Altitude', 'Altitud media'), (252, 'Extreme Altitude', 'Altitud extrema')]:
        v = number(2, density, {'bucket': name}, 'mean', 3)
        bar(p, label, val(v), v, y, 1.1, RED if name == 'Extreme Altitude' else MUTED)
    p.text('Densidad relativa: menor valor significa aire menos denso.', 40, 306, 640, 12, MUTED)
    end()

    # 3 · the physics
    observed = 'reports/physics/delta_extreme_vs_no.csv'
    elastic = 'reports/physics_v2/magnus_elasticity.csv'
    obs = number(3, observed, {'pitch_family': 'Four-Seam', 'outcome': 'ivb'}, 'delta')
    cf = number(3, shape, {'shape': '4S'}, 'delta_ivb_in')
    el = number(3, elastic, {'pitch_family': 'Four-Seam'}, 'elasticity')
    p = slide('La recta pierde subida: el Magnus cae\nmás rápido que el aire', observed + '; ' + shape + ' · 4S; ' + elastic + '.',
              '55 s. Observado dentro del mismo lanzador contra el cambio del modelo para la forma 4S. La elasticidad dice cuánto cae la fuerza de Magnus con la densidad: la física de libro dice 1, medimos más.')
    for x, value, label in [(40, obs+'"', 'Observado\naltitud extrema − sin altitud'),
                            (262, cf+'"', 'Modelo, forma 4S\nCDMX − nivel del mar'),
                            (484, el, 'Elasticidad del Magnus\nde la recta (libro: 1)')]:
        card(p, x, 122, 196, 166)
        stat(p, x, 122, 196, value, label)
    p.text('IVB: movimiento vertical inducido, pulgadas. El simulador usa la elasticidad medida.', 40, 306, 640, 12, MUTED)
    end()

    # 4 · validation against the rubric's null
    hold = 'reports/validation/headline.csv'
    skill = 'reports/stuff_v2/skill.csv'
    level = 'reports/stuff_v2/pitcher_level.csv'
    pk = number(4, hold, {'split': 'pitcher_kfold', 'model': 'Pitching+'}, 'skill', 2, 100)
    v2 = number(4, skill, {'model': 'Stuff+ v2', 'split': 'pitcher_kfold'}, 'skill_vs_baseline', 2, 100)
    r50 = number(4, level, {'model': 'Stuff+ v2 (familia)'}, 'first50_r', 3)
    r50obs = number(4, level, {'model': 'Stuff+ v2 (familia)'}, 'observed_first50_r', 3)
    p = slide('Le ganamos al modelo nulo de la rúbrica', hold + '; ' + skill + '; ' + level + '.',
              '75 s. Skill: fracción del RMSE que el modelo le quita al nulo por conteo, por pitcher. Stuff+ v2 deja que el conteo valore los resultados y ya le gana al nulo. A nivel pitcher, 50 lanzamientos de Stuff+ dicen más del resto de la temporada que los resultados del propio pitcher.')
    bar(p, 'Pitching+', val(pk), pk+' %', 120, 5)
    bar(p, 'Stuff+ v2', val(v2), '+'+v2+' %', 152, 5)
    p.text('% de RMSE que baja frente al nulo por conteo, validación por pitcher.', 40, 184, 640, 11, MUTED)
    card(p, 40, 210, 316, 120)
    stat(p, 40, 210, 316, r50, 'r de Stuff+ v2 con 50 lanzamientos contra\nel resto de la temporada', RED, 34, 12)
    card(p, 364, 210, 316, 120)
    stat(p, 364, 210, 316, r50obs, 'r de los resultados del propio pitcher\ncon los mismos 50 lanzamientos', MUTED, 34, 12)
    end()

    # 5 · physics fixes the altitude bias
    bias = 'reports/stuff_v2/altitude_bias.csv'
    b0 = number(5, bias, {'model': 'baseline (conteo)'}, 'bias_obs_minus_pred', 4)
    b1 = number(5, bias, {'model': 'Stuff+ v1'}, 'bias_obs_minus_pred', 4)
    b2 = number(5, bias, {'model': 'Stuff+ v2'}, 'bias_obs_minus_pred', 4)
    cut = derive(5, '1-a/b', [(bias, {'model': 'Stuff+ v2'}, 'bias_obs_minus_pred'),
                              (bias, {'model': 'baseline (conteo)'}, 'bias_obs_minus_pred')], 0)
    p = slide('La física corrige el sesgo de altitud', bias + ' · entrenado sin altitud extrema, evaluado en ella.',
              '55 s. Se entrena sin un solo juego de altitud extrema y se predice esa altitud. El sesgo es carreras observadas menos predichas por lanzamiento: positivo = el modelo subestima el daño. Con la densidad del aire, Stuff+ v2 casi lo elimina.')
    bar(p, 'Nulo (conteo)', val(b0), '+'+b0, 126, 0.012, MUTED)
    bar(p, 'Stuff+ v1', val(b1), '+'+b1, 164, 0.012)
    bar(p, 'Stuff+ v2', val(b2), '+'+b2, 202, 0.012)
    p.rect(40, 246, 640, 44, PANEL, 6)
    p.rect(40, 246, 4, 44, RED)
    p.text(f'Stuff+ v2 recorta el sesgo {cut} % frente al nulo.', 58, 259, 610, 17, TEXT, True)
    p.text('Carreras por lanzamiento que el modelo subestima en un aire que no vio.', 40, 306, 640, 12, MUTED)
    end()

    # 6 · shapes
    p = slide('Las pérdidas dependen de la forma del pitch', shape + ' · delta_stuff; misma forma, CDMX menos nivel del mar.',
              '65 s. Ninguna forma gana en términos absolutos. La curva 12-6 y el sweeper pierden menos que la liga; la recta de mucho ride pierde más. Hipótesis operativa para reclutar: no pagar por ride de recta.')
    for y, label, key in [(122, 'Curva 12-6', '12-6 CB'), (160, 'Sweeper', 'sweeper'), (198, 'Recta 4S', '4S'), (236, 'Recta de ride', 'riding 4S')]:
        v = number(6, shape, {'shape': key}, 'delta_stuff', 1)
        bar(p, label, val(v), v, y, 19, RED if key == 'riding 4S' else MUTED)
    p.text('Longitud = pérdida de Stuff+. Ninguna forma gana en absoluto.', 40, 306, 640, 12, MUTED)
    end()

    # 7 · staff
    saved = 'reports/staff_audit/runs_saved.csv'

    def saved_num(franchise, component, column, digits=1):
        return number(7, saved, {'franchise': franchise, 'component': component}, column, digits)
    r5 = saved_num('5', 'total', 'runs_per_season'); w5 = saved_num('5', 'total', 'wins_per_season', 2)
    r11 = saved_num('11', 'total', 'runs_per_season'); w11 = saved_num('11', 'total', 'wins_per_season', 2)
    m5 = saved_num('5', 'mezcla de lanzamientos en casa', 'runs_per_season')
    a5 = saved_num('5', 'asignación de aperturas en casa', 'runs_per_season')
    m11 = saved_num('11', 'mezcla de lanzamientos en casa', 'runs_per_season')
    a11 = saved_num('11', 'asignación de aperturas en casa', 'runs_per_season')
    p = slide('Dos sedes candidatas, dos cuentas de carreras', saved + ' · franquicias 5 y 11, por separado.',
              '55 s. Identificación inferida: son las dos franquicias de aire más ligero, no un staff confirmado de Diablos. Carreras del modelo: mezcla de lanzamientos en casa más reasignación de aperturas. No sumar 5 y 11.')
    for x, name, r, w, m, a in [(40, 'Franquicia 5', r5, w5, m5, a5), (364, 'Franquicia 11', r11, w11, m11, a11)]:
        card(p, x, 118, 316, 158)
        p.text(name, x+18, 132, 280, 14, MUTED, True)
        p.text(r+' carreras', x+18, 156, 280, 32, RED, True)
        p.text(w+' victorias por temporada', x+18, 202, 280, 15, TEXT)
        p.text('Mezcla '+m+' · aperturas '+a, x+18, 236, 280, 13, MUTED)
    p.text('Identificación inferida. No se suman.\nLa sede sigue sin confirmarse.', 40, 290, 640, 13, TEXT, leading=17)
    end()

    # 8 · the product
    p = slide('El producto: dos aires, un lanzamiento', 'Capturas del simulador 3D y del explicador de física (Playwright, web/e2e/screenshots, HACK-051).',
              '50 s. Simulador 3D: el mismo release a nivel del mar y en el aire del Harp Helú, con el motor de vuelo que usa la elasticidad medida. El explicador de física deja mover la densidad y ver la subida perdida. Las cifras de la app salen de los mismos agregados.')
    p.rect(38, 94, 312, 182, PANEL, 4)
    p.image(ASSETS/'simulador_transmision.jpg', 40, 96, 308, 178)
    p.rect(370, 94, 312, 202, PANEL, 4)
    p.image(ASSETS/'fisica.jpg', 372, 96, 308, 198)
    p.text('Simulador 3D: dos aires, mismo release.', 40, 286, 308, 11, MUTED)
    p.text('Física: arrastre y Magnus contra la densidad.', 372, 306, 308, 11, MUTED)
    end()

    # 9 · value
    value = 'reports/business_value/runs_per_win.csv'
    gm = 'reports/business_value/gm_example.csv'
    rpw_l = number(9, value, {'entorno': 'liga'}, 'carreras_por_victoria', 1)
    rpw_x = number(9, value, {'entorno': 'Extreme Altitude'}, 'carreras_por_victoria', 1)
    wins = number(9, gm, {'escala': 'Stuff+', 'entorno': 'Harp Helú (bucket Extreme)'}, 'victorias', 2)
    runs = number(9, gm, {'escala': 'Stuff+', 'entorno': 'Harp Helú (bucket Extreme)'}, 'carreras', 1)
    p = slide('El entorno cambia el valor de una carrera', value + '; ' + gm + ' · puntos de Stuff+ v2 fuera de pliegue.',
              '65 s. En altura se anota más, así que cada carrera ahorrada vale menos victorias. +10 puntos de Stuff+ sostenidos en 2,500 lanzamientos es un escenario de orden de magnitud, no el ahorro del staff de la lámina 7.')
    for x, v, label, color in [(40, rpw_l, 'Carreras por victoria\nliga', MUTED),
                               (262, rpw_x, 'Carreras por victoria\naltitud extrema', RED),
                               (484, wins, 'Victorias en el Harp Helú\npor +10 de Stuff+', RED)]:
        card(p, x, 118, 196, 146)
        stat(p, x, 118, 196, v, label, color, 38)
    p.text(f'+10 puntos en 2,500 lanzamientos = {runs} carreras (escenario).', 40, 280, 640, 15, TEXT, True)
    p.text('El ahorro del staff está en la lámina 7, por franquicia.', 40, 306, 640, 12, MUTED)
    end()

    # 10 · demo
    p = slide('La demo compara el mismo arsenal', 'Producto web en el repo. Respaldo: este PDF y los CSV agregados de reports/.',
              '70 s. Abrir el simulador con una forma y dos densidades, pasar al explicador de física y a la página de altitud. Si la red falla, quedarse en este PDF. Cerrar: las franquicias 5 y 11 no se suman y la sede no está confirmada.')
    p.rect(378, 94, 304, 177, PANEL, 4)
    p.image(ASSETS/'simulador_receptor.jpg', 380, 96, 300, 173)
    for i, line in enumerate(['Simulador: la misma forma, dos aires.', 'Física: la densidad en vivo.', 'Altitud: los 20 parques por densidad.']):
        p.rect(40, 104+i*52, 4, 36, RED)
        p.text(line, 56, 110+i*52, 316, 15, TEXT, True)
    p.text('Respaldo offline: este PDF y los CSV agregados.\nLa sede del staff sigue sin confirmarse.', 40, 288, 640, 13, MUTED, leading=17)
    end()
    pdf.save(); prs.save(OUT/'deck.pptx')

    # One-page executive brief, designed for printing (white page).
    c = Canvas(str(ONE/'one-pager.pdf'), pagesize=(595, 842))
    c.setTitle('Stuff+ 2240 · resumen ejecutivo v2')
    one_tex = []
    p = Page(pdf=c, width=595, height=842, tex=one_tex, tex_base=ONE)
    p.rect(0, 0, 595, 10, RED)
    p.rect(38, 40, 8, 8, RED)
    p.text('STUFF+ 2240   /   DIABLOS ROJOS', 54, 37, 463, 9.5, GRAY, True)
    p.text('El arsenal pierde valor en CDMX', 38, 64, 519, 26, INK, True)
    p.text('Resumen ejecutivo v2 para el GM y el cuerpo técnico', 38, 102, 519, 12.5, GRAY)
    for x, v, label in [(38, league, 'ΔStuff+ de la liga\nen el Harp Helú'),
                        (213, cut + ' %', 'del sesgo de altitud que\nrecorta Stuff+ v2'),
                        (388, r5 + ' / ' + r11, 'carreras por temporada,\nfranquicias 5 / 11')]:
        p.rect(x, 132, 169, 82, LIGHT, 5)
        p.text(v, x+12, 142, 150, 24, RED, True)
        p.text(label, x+12, 176, 150, 9.5, INK, leading=12.5)

    def section(title, y):
        p.rect(38, y+2, 3, 15, RED)
        return p.text(title, 48, y, 509, 14, INK, True) + 6

    def body(text, y):
        return p.text(text, 38, y, 555, 11, INK, leading=15.5) + 7  # Fira Sans is narrower than the DejaVu used to wrap
    y = section('Hallazgos', 236)
    y = body(f'El mismo lanzamiento pierde {league.replace("−", "")} puntos de Stuff+ en promedio en el aire del Harp Helú. La recta de mucho ride es la que más pierde; la curva 12-6 y el sweeper, las que mejor aguantan. La fuerza de Magnus de la recta cae con la densidad elevada a {el}, más que la física de libro.', y)
    y = body(f'Stuff+ v2 le gana al modelo nulo de la rúbrica (+{v2} % de RMSE por pitcher; Pitching+ {pk} %). Con 50 lanzamientos predice el resto de la temporada con r = {r50}, contra {r50obs} de los propios resultados del pitcher.', y)
    y = body(f'Entrenado sin altitud extrema, Stuff+ v2 recorta {cut} % el daño que el nulo subestima allí: la física extrapola a un aire que no vio.', y)
    y = section('Decisión propuesta', y+10)
    y = body('No pagar por ride de recta; preferir secundarios que aguantan el aire (curva 12-6, sweeper). En casa, mover uso hacia el lanzamiento que mejor resiste y dar las aperturas de casa al abridor más resistente. Confirmar antes con el club cuál franquicia es la sede.', y)
    y = section('Valor', y+10)
    y = body(f'Franquicia 5: {r5} carreras ({w5} victorias) por temporada. Franquicia 11: {r11} carreras ({w11} victorias). Identificación inferida: no se suman. Una carrera vale {rpw_l} por victoria en la liga y {rpw_x} en altitud extrema; +10 de Stuff+ en 2,500 lanzamientos son {wins} victorias en el Harp Helú.', y)
    assert y < 720, y
    p.rect(38, 734, 519, .8, LIGHT)
    p.text('Fuentes y trazabilidad', 38, 744, 519, 10.5, INK, True)
    p.text('altitude_delta/delta_by_shape.csv; physics_v2/magnus_elasticity.csv; stuff_v2/skill.csv,\npitcher_level.csv y altitude_bias.csv; validation/headline.csv; staff_audit/runs_saved.csv;\nbusiness_value/runs_per_win.csv y gm_example.csv. Todos bajo reports/. Filas, columnas\ny hashes: demo/deck/sources.json.', 38, 762, 519, 8.5, GRAY, leading=12)
    c.showPage(); c.save()
    typeset(tex_pages, one_tex)

    (OUT/'sources.json').write_text(json.dumps({'demo_data': False, 'exploratory': True, 'model': 'Stuff+ v2 (HACK-057)',
        'source_sha256': sources, 'evidence': evidence, 'derived': derived}, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    timing = [45, 55, 55, 75, 55, 65, 55, 50, 65, 70]
    rows = ['# Presentación ejecutiva v2 · HACK-036, HACK-059', '',
            '10 diapositivas para una entrevista de 10 minutos, con Stuff+ v2 en producción (HACK-057). Fondo carbón, rojo #D71920. El PDF se compone en LaTeX (`deck.tex`, Beamer + TikZ, Fira Sans) con Tectonic; sin Tectonic se queda la versión de ReportLab con DejaVu Sans de `fonts/` (licencia en `fonts/LICENSE_DEJAVU`). Las dos usan las mismas líneas que el PPTX editable (Arial).', '',
            '**Estado:** evidencia oficial agregada. Las franquicias 5 y 11 no se suman: identificación inferida y sede sin confirmar.', '',
            '| Lámina | Idea principal | Tiempo | Fuente de cifras |', '|---|---|---:|---|']
    for i, (t, f, sec) in enumerate(zip(titles, footers, timing), 1):
        rows.append(f'| {i} | {t.replace(chr(10), " ")} | {sec} s | {f} |')
    rows += ['', '## Archivos y reconstrucción', '',
             '`deck.pptx`: editable. `deck.pdf`: versión de presentación. `../one-pager/one-pager.pdf`: una página para el GM. `assets/`: capturas del producto tomadas por las pruebas e2e de HACK-051 (`web/e2e/screenshots/`).', '',
             '`sources.json` guarda archivo, fila, columna, formato y SHA-256 de cada cifra, y para las dos cifras derivadas (promedio de la liga y recorte del sesgo) sus entradas y la fórmula. Las notas de cada diapositiva llevan el guion oral.', '',
             'Reconstruir con Python 3.12, `demo/deck/requirements.txt` y Tectonic (`scoop install tectonic` o `brew install tectonic`): `python demo/deck/build.py`. Validar: `python demo/deck/verify.py`.', '',
             '## Revisión antes de entrevista', '',
             'Confirmar con el club cuál de las franquicias 5 y 11 es la sede. La conversión a victorias usa los puntos de Stuff+ v2 fuera de pliegue (`reports/business_value/stuff_scale_calibration.csv`, HACK-078).', '',
             '## Verificación', '',
             'Verify del backlog: `test -s demo/deck/deck.pdf && test -s demo/one-pager/one-pager.pdf`. Trazabilidad: `python demo/deck/verify.py`.']
    (OUT/'README.md').write_text('\n'.join(rows)+'\n', encoding='utf-8')
    (ONE/'README.md').write_text('# Resumen ejecutivo v2\n\n`one-pager.pdf`: una página en español para el GM, con Stuff+ v2. Fuentes en `../deck/README.md` y `../deck/sources.json`. Reconstrucción: `python demo/deck/build.py`.\n', encoding='utf-8')
    print(f'Built 10-slide PPTX + presentation PDF + one-page PDF with {len(evidence)} source bindings and {len(derived)} derived values.')


if __name__ == '__main__': main()
